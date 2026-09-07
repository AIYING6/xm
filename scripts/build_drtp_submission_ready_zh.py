# -*- coding: utf-8 -*-
"""Build a Chinese journal-style DRTP manuscript from frozen local evidence.

This script is deliberately evidence-bound: only the completed formal five-seed
comparison is quantified; incomplete or non-confirmatory material is kept as a
clearly labelled research boundary.
"""
from __future__ import annotations

from pathlib import Path
import csv
import math
from collections import defaultdict

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "drtp_submission_ready"
FIG = OUT / "figures_submission_ready"
DATA = ROOT / "output" / "drtp_relay_failure_anonymous_reproducibility_v5" / "source_data"
FORMAL_REPORT = DATA / "formal_2301_2305" / "DRTP_UTR_Q2_FORMAL_FIVE_SEED_CONFIRMATION_REPORT.md"
RAW = DATA / "formal_2301_2305" / "evaluations" / "final_10m" / "raw_episode_metrics.csv"


def cn_font(run, name="宋体", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_border(cell, color="B7C9D6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    cn_font(r, "宋体", 9.5)


def add_para(doc, text, first=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.first_line_indent = Cm(0.74) if first else Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    cn_font(r, "宋体", 10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    cn_font(r, "黑体" if level == 1 else "宋体", 14 if level == 1 else 11.5, True, (26, 76, 114) if level == 1 else (0,0,0))
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "1A4C72")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h)); cn_font(r, "宋体", 8.8, True, (255,255,255))
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            if row_i % 2 == 1: set_cell_shading(cells[i], "EEF4F8")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value)); cn_font(r, "宋体", 8.7)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths): cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def load_rows():
    with RAW.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def fields(rows):
    return list(rows[0].keys()) if rows else []


def guess_key(names, options):
    lower = {n.lower():n for n in names}
    for o in options:
        if o.lower() in lower: return lower[o.lower()]
    for n in names:
        if any(o.lower() in n.lower() for o in options): return n
    return None


def make_figures():
    FIG.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.sans-serif":["Microsoft YaHei", "SimHei", "Arial Unicode MS"], "axes.unicode_minus":False, "font.size":9})
    navy, blue, orange, pale = "#18496B", "#4B89B8", "#E67E22", "#EAF2F8"
    # Fig 1: problem mechanism
    fig, ax = plt.subplots(figsize=(8.0, 3.3)); ax.set_xlim(0, 10); ax.set_ylim(0, 4); ax.axis("off")
    nodes=[(1.0,2.8,"侦察 UAV",blue),(4.0,2.8,"继电 UAV",orange),(7.0,2.8,"攻击 UAV",blue),(7.0,0.7,"任务完成",navy)]
    for x,y,t,c in nodes:
        circ=plt.Circle((x,y),0.48,facecolor=c,edgecolor="white",lw=1.8); ax.add_patch(circ); ax.text(x,y,t,ha="center",va="center",color="white",weight="bold")
    ax.annotate("目标信息",(3.45,2.8),(1.55,2.8),arrowprops=dict(arrowstyle="->",lw=2,color=navy),ha="center",va="bottom")
    ax.annotate("合法缓存 / 转发",(6.45,2.8),(4.55,2.8),arrowprops=dict(arrowstyle="->",lw=2,color=navy),ha="center",va="bottom")
    ax.annotate("协同拦截",(7,1.25),(7,2.3),arrowprops=dict(arrowstyle="->",lw=2,color=navy),ha="left")
    ax.plot([4,7],[2.8,2.8],color="#C0392B",lw=4,alpha=.85); ax.text(5.5,3.15,"继电故障：信息支撑路径退化",color="#A93226",ha="center",weight="bold")
    ax.text(5.5,0.15,"研究问题：在合法信息边界固定时，如何通过训练暴露塑造故障下的协同策略？",ha="center",color=navy,weight="bold")
    fig.tight_layout(); fig.savefig(FIG/"fig1_topology_degradation.png",dpi=300,bbox_inches="tight"); plt.close(fig)
    # Fig 2: framework
    fig, ax = plt.subplots(figsize=(8.0, 4)); ax.set_xlim(0, 10); ax.set_ylim(0, 5); ax.axis("off")
    boxes=[(0.3,3.4,2.0,0.9,"冻结故障条件库\nN / F0 / TE / TL / DS / DL / CP",pale), (3.0,3.4,1.7,0.9,"组级回报\n难度代理", "#FDF2E9"),(5.5,3.4,1.7,0.9,"有界概率\nq 更新", "#E8F6F3"),(7.8,3.4,1.7,0.9,"Reset 条件\n暴露分配", "#EAF2F8"),(3.8,1.0,2.5,1.0,"共享 MAPPO 训练\nActor / Critic / PPO / 奖励不变", "#D6EAF8")]
    for x,y,w,h,t,c in boxes:
        ax.add_patch(plt.Rectangle((x,y),w,h,facecolor=c,edgecolor=navy,lw=1.4)); ax.text(x+w/2,y+h/2,t,ha="center",va="center",weight="bold",color=navy)
    for a,b in [((2.3,3.85),(3,3.85)),((4.7,3.85),(5.5,3.85)),((7.2,3.85),(7.8,3.85)),((8.65,3.4),(6.3,2.0)),((3.8,2.0),(1.3,3.4))]:
        ax.annotate("",b,a,arrowprops=dict(arrowstyle="->",lw=1.8,color=navy))
    ax.text(5,4.72,"DRTP：训练期拓扑感知暴露塑形",ha="center",fontsize=13,weight="bold",color=navy)
    ax.text(5,.35,"唯一干预：故障组的 reset 采样质量；执行期信息、网络容量与优化目标均保持固定。",ha="center",color="#4D5656")
    fig.tight_layout(); fig.savefig(FIG/"fig2_drtp_framework.png",dpi=300,bbox_inches="tight"); plt.close(fig)
    # Fig 3: experimental control
    fig, ax = plt.subplots(figsize=(8.0,3.2)); ax.axis("off"); ax.set_xlim(0,10); ax.set_ylim(0,4)
    for x,label,c in [(0.5,"UTR\n故障组均匀暴露", "#D6EAF8"),(4.0,"共同项\n3DOF 环境 / MAPPO / 奖励 / 预算 / 评价带", "#F4F6F7"),(7.7,"DRTP\n有界自适应暴露", "#D5F5E3")]:
        w=2.1 if x!=4 else 2.8; ax.add_patch(plt.Rectangle((x,1.15),w,1.7,facecolor=c,edgecolor=navy,lw=1.5)); ax.text(x+w/2,2,label,ha="center",va="center",weight="bold",color=navy)
    ax.annotate("",(4.0,2),(2.6,2),arrowprops=dict(arrowstyle="->",lw=2,color=navy)); ax.annotate("",(7.7,2),(6.8,2),arrowprops=dict(arrowstyle="->",lw=2,color=navy))
    ax.text(5,.4,"主要因果对照：除训练 reset 阶段的故障组分配外，其余变量保持一致。",ha="center",color=navy,weight="bold")
    fig.tight_layout(); fig.savefig(FIG/"fig3_controlled_comparison.png",dpi=300,bbox_inches="tight"); plt.close(fig)
    # Fig 4: formal endpoint means. Values independently frozen in report.
    endpoints=["名义", "F0", "非名义平均", "非名义最差"]
    utr=[174.30,144.64,144.70,124.44]; drtp=[207.78,196.77,199.70,187.45]
    x=range(4); fig, ax=plt.subplots(figsize=(7.6,3.7)); width=.34
    ax.bar([i-width/2 for i in x],utr,width,label="UTR",color="#AAB7B8")
    ax.bar([i+width/2 for i in x],drtp,width,label="DRTP",color=blue)
    for i,(u,d) in enumerate(zip(utr,drtp)): ax.text(i+width/2,d+3,f"+{d-u:.2f}",ha="center",fontsize=8,color=navy,weight="bold")
    ax.set_ylabel("固定终点评价回报 J"); ax.set_xticks(list(x),endpoints); ax.spines[["top","right"]].set_visible(False); ax.legend(frameon=False,ncol=2); ax.grid(axis="y",alpha=.22)
    fig.tight_layout(); fig.savefig(FIG/"fig4_formal_results.png",dpi=300,bbox_inches="tight"); plt.close(fig)
    # Fig 5: q mechanism as a strict schematic, no quantitative q data inferred.
    fig, ax=plt.subplots(figsize=(7.8,3.8)); groups=["F0","TE","TL","DS","DL","CP"]; u=[1/6]*6; d=[.11,.18,.13,.20,.16,.22]
    x=range(6); ax.bar([i-.18 for i in x],u,.36,label="UTR：固定均匀",color="#BDC3C7"); ax.bar([i+.18 for i in x],d,.36,label="DRTP：示意性有界重分配",color=blue)
    ax.set_ylim(0,.4); ax.set_xticks(list(x),groups); ax.set_ylabel("故障组条件质量 q"); ax.legend(frameon=False,ncol=2); ax.grid(axis="y",alpha=.2); ax.spines[["top","right"]].set_visible(False)
    ax.text(.02,.95,"说明：本图呈现机制结构，不表示任何单一训练种子的实测最终 q。",transform=ax.transAxes,va="top",fontsize=8,color="#7B241C")
    fig.tight_layout(); fig.savefig(FIG/"fig5_mechanism_schematic.png",dpi=300,bbox_inches="tight"); plt.close(fig)


def build_doc():
    make_figures()
    OUT.mkdir(parents=True, exist_ok=True)
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.1); sec.left_margin=Cm(2.35); sec.right_margin=Cm(2.35)
    styles=doc.styles
    styles["Normal"].font.name="宋体"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体"); styles["Normal"].font.size=Pt(10.5)
    # title page
    for _ in range(4): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
    r=p.add_run("面向通信拓扑退化的三无人机协同决策训练暴露塑形方法"); cn_font(r,"黑体",17.5,True,(20,61,91))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("——基于有界自适应故障组重加权的受控研究"); cn_font(r,"宋体",13,False,(76,86,96))
    for _ in range(5): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("中文投稿初稿  V2"); cn_font(r,"宋体",12,False,(76,86,96))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("2026 年 9 月"); cn_font(r,"宋体",12)
    doc.add_page_break()
    add_heading(doc,"摘  要",1)
    abstract=("多无人机协同任务依赖目标信息沿感知、通信与任务支撑路径到达相应角色。继电节点故障不仅减少可用通信边，还会改变攻击角色可合法使用的信息来源，从而使训练阶段接触的故障分布与部署阶段的拓扑退化需求发生失配。本文研究训练暴露分配这一可隔离变量：在策略骨干、集中训练分散执行框架、PPO 目标、奖励、观测、动作、执行期信息边界、名义条件质量、故障支持集合和训练预算保持一致时，改变多个冻结故障组在 reset 阶段的采样质量，能否改善故障条件下的协同任务端点。为此，提出一种有界自适应拓扑重加权方法 DRTP。该方法以故障组回报相对名义组回报的差异作为训练期困难代理，周期性更新故障组采样分布，并以概率上下界保证所有冻结组持续被暴露。\n\n在轻量三自由度、三架异构无人机协作拦截环境的正式五种子受控队列中，DRTP 相比均匀拓扑随机化 UTR 的 F0、非名义平均和非名义最差回报的配对平均增益分别为 52.13、55.00 和 63.01，三项端点均为 5/5 个训练种子正向；非名义平均超时率由 0.874 降至 0.694。碰撞率由 0.0051 升至 0.0080，因此本文不将任务得分提升表述为全面安全改进。采样遥测表明训练暴露分布确已改变，但不足以单独证明策略内部的信息恢复机制。本文结论限定于冻结的三无人机继电故障合同：拓扑语义驱动的有界暴露重加权可在该受控设置中提升队列级故障任务表现；跨训练队列的一致性、修复后的跨规模结果以及语义消融仍需作为后续证据补充。")
    add_para(doc,abstract)
    p=doc.add_paragraph(); r=p.add_run("关键词："); cn_font(r,"黑体",10.5,True); r=p.add_run("多智能体强化学习；无人机协同；通信拓扑退化；继电故障；训练暴露分配；固定终点评估"); cn_font(r,"宋体",10.5)
    # Intro
    add_heading(doc,"1 引言",1)
    for t in [
        "在多无人机协同任务中，飞行器的运动能力并不能单独决定任务是否完成。侦察、继电与攻击等角色必须在局部感知、通信约束和任务时序共同限制下形成可执行的协作链。当信息不能沿合法路径及时抵达需要使用它的角色时，个体即使具有足够的机动能力，也可能无法完成后续协同动作。因此，通信拓扑退化应被理解为改变协同决策信息结构的事件，而不只是叠加在动力学上的随机噪声。",
        "现有无人机协同决策研究已从任务建模、策略学习与仿真验证等层面讨论了通信约束和部分可观测性的重要性[1-4]。另一方面，强化学习中的课程学习、域随机化和优先级采样表明，训练条件的组织方式会影响策略所获得的能力[2,5]。然而，在继电依赖的协作任务中，不同故障条件并非只在扰动幅度上不同：它们可能改变故障起始时刻、持续时间、缓存时效以及可用信息支撑路径。若训练预算在这些条件之间始终均匀分配，训练过程虽然覆盖了故障多样性，却未必把更多学习机会分配给当前策略最难适应的结构性退化条件。",
        "本文不试图在执行期恢复不可见状态，也不引入新的通信协议、图网络输入或奖励塑形。相反，我们将问题缩小为一个可检验的训练干预：在执行接口与学习器完全匹配的前提下，仅改变训练 reset 时冻结故障组的暴露质量，是否足以改变固定终点评价下的故障协同表现。该定位使拓扑退化问题、算法变化项与主要因果对照可以被清晰分离。",
        "据此，本文提出动态继电拓扑重加权方法 DRTP。其核心不是为每个故障组预设一套固定难度标签，而是以组级已完成 episode 回报相对名义组的偏离构建训练期困难代理，在给定上下界内动态调整故障组质量。该设计保持名义条件质量和故障支持集合不变，因而将训练暴露分配作为唯一主要差异。本文的主要贡献包括：\n（1）将继电故障下的通信拓扑退化表述为合法信息支撑路径变化，并据此提出一个受控的训练分布设计问题；\n（2）提出不改变策略结构、优化目标和执行期信息边界的有界自适应故障组重加权机制；\n（3）以固定 10M 终点、五个训练种子、共同评价带和可靠性指标为基础，报告主要受控比较及其跨队列适用范围。"
    ]: add_para(doc,t)
    doc.add_picture(str(FIG/"fig1_topology_degradation.png"),width=Cm(15.7)); add_caption(doc,"图 1 继电故障下的信息支撑路径退化与协同决策问题示意")
    # Related work
    add_heading(doc,"2 相关研究与本文定位",1)
    add_heading(doc,"2.1 通信受限无人机协同与拓扑退化",2)
    add_para(doc,"无人机集群协同通常需要在局部观测、异构角色和通信约束下完成感知—信息传递—决策—行动闭环。有人—无人机协同、无人集群博弈对抗和空战智能决策等研究均指出，信息共享条件会直接影响任务分工与协同效率[1,3,4,6]。本文沿用这一系统视角，但将研究焦点置于训练阶段：继电节点故障改变合法信息路径后，不同拓扑条件应获得怎样的训练暴露。")
    add_heading(doc,"2.2 训练条件组织与优先级机制",2)
    add_para(doc,"课程学习通过任务阶段或难度递进改善学习过程，优先级机制则依据通用学习信号改变采样频率。两类思想说明训练分布本身是可设计变量。DRTP 与它们的关系是：DRTP 同样在训练阶段调整条件暴露，但其采样对象是预先冻结的故障拓扑组，适应信号为相对名义任务的组级表现，且采样概率受到显式边界约束。因而，本文不把外部优先级方法当作与 UTR 完全等价的因果对照，而把 UTR–DRTP 的同架构匹配比较视为主要证据。")
    add_table(doc,["维度","UTR","通用优先级采样","DRTP"],[
        ["采样对象","冻结故障组","rollout/样本或环境片段","冻结故障拓扑组"],
        ["适应信号","无","通用优先级信号","组级名义相对回报"],
        ["拓扑语义","无显式使用","不要求","作为条件组定义"],
        ["执行期接口","不变","视实现而定","不变"],
        ["主要作用位置","reset","训练采样","reset"]], [3.1,3.6,4.7])
    add_caption(doc,"表 1 训练条件组织方法的概念性比较")
    # Task and method
    add_heading(doc,"3 问题建模与 DRTP 方法",1)
    add_heading(doc,"3.1 三无人机继电依赖协作任务",2)
    add_para(doc,"环境包含侦察、继电和攻击三类蓝方 UAV 以及一个目标，采用轻量三自由度离散控制。侦察角色获得目标线索；在攻击角色尚未进入直接感知包络时，攻击角色只能基于环境规则允许的、仍然新鲜的本地或送达缓存信息行动。继电节点在故障窗口内不能收发通信，故障后由侦察到攻击的直接送达只有在满足既定时序与合法性条件时才能成为替代来源。这里的“路径重构”指环境认可的信息来源改变，不表示策略访问了隐藏全局状态。")
    add_heading(doc,"3.2 冻结条件组与训练暴露",2)
    add_para(doc,"每个训练条件写作 c=(g,m,τ,δ)，其中 g 为条件组，m 为组内冻结成员，τ 与 δ 分别为故障起始时刻和持续时间。训练集包含名义组 N 与六个故障组 F0、TE、TL、DS、DL 和 CP。名义质量固定为 0.50，剩余 0.50 在故障组之间分配；组内成员均匀抽取。UTR 的故障组条件质量为 q_u(g)=1/6。DRTP 从相同初值开始，仅在训练过程中更新 q(g)。")
    add_heading(doc,"3.3 名义相对困难代理与有界更新",2)
    add_para(doc,"设 R̄g 和 R̄N 分别为故障组 g 与名义组的完成 episode 回报指数滑动平均。DRTP 使用 d_g=clip[(R̄N−R̄g)/max(|R̄N|,ε),0,dmax] 构造困难代理。该量刻画当前训练轨迹中故障组相对名义组的任务落差；它不是图连通性、最短路长度、策略不确定性或因果难度的估计。warm-up 后，采样器按照 q̃_g∝q_g exp{η(d_g−d̄)} 形成候选分布，再经平滑与有界单纯形投影得到 q_{u+1}。本文冻结的下界和上界分别为 0.05 与 0.35，从而防止单一故障组垄断训练暴露。")
    add_para(doc,"DRTP 的唯一作用位置是 reset 阶段。actor、critic、PPO 目标、奖励、观测、动作、环境动力学、名义条件质量、故障支持集合和终点评价协议均与 UTR 相同。因此，若固定终点出现差异，其直接可归因的干预是训练期间故障组的暴露分配，而不是策略容量或执行期额外信息。")
    doc.add_picture(str(FIG/"fig2_drtp_framework.png"),width=Cm(15.7)); add_caption(doc,"图 2 DRTP 训练期拓扑感知暴露塑形框架")
    doc.add_picture(str(FIG/"fig3_controlled_comparison.png"),width=Cm(15.7)); add_caption(doc,"图 3 UTR–DRTP 主要受控比较的固定项与唯一变化项")
    # Experiments
    add_heading(doc,"4 实验设计",1)
    add_heading(doc,"4.1 研究问题与报告原则",2)
    add_para(doc,"本文围绕以下问题组织证据。RQ1：DRTP 是否在冻结正式合同内改善典型故障与跨条件故障端点？RQ2：回报变化是否伴随超时、碰撞、约束违规或故障触发方面的可靠性权衡？RQ3：采样器是否确实改变了训练暴露，以及这些遥测能证明什么？RQ4：独立队列和外部参考如何限定结论适用范围？训练种子是独立统计单位；同一策略在不同 episode 和条件上的评价记录用于估计该种子的端点，不被当作额外独立训练重复。")
    add_table(doc,["项目","冻结设置"],[
        ["任务环境","3DOF 三 UAV 异构协作；继电依赖信息边界"],
        ["主要比较","UTR-SG 与 DRTP-SG"],
        ["唯一变化","故障组 reset 条件质量：均匀 vs 有界自适应"],
        ["共同项","Actor/Critic、PPO、奖励、观测、动作、环境、预算、评价带"],
        ["训练单位","5 个配对训练种子"],
        ["训练预算","39,063 updates；10,000,128 环境步/条轨迹"],
        ["检查点规则","固定最终 10M checkpoint"],
        ["评价带","12 条固定条件，每条件 100 episodes"],
    ],[3.3,8.1]); add_caption(doc,"表 2 正式受控实验协议")
    add_heading(doc,"4.2 端点与证据分层",2)
    add_para(doc,"主要任务端点为名义回报、F0 回报、非名义平均回报与非名义最差回报；可靠性指标为成功、超时、碰撞与约束违规。由于正式评价条件属于冻结训练支持集合，本文将非名义端点称为跨条件故障端点，而不称为严格 OOD。正式五种子配对队列承担主要比较；后续独立队列用于考察跨队列方向；PLR-style 外部参考用于定位通用优先级机制的相对表现。不同层级不合并为一个扩大后的确认性样本。")
    # Results
    add_heading(doc,"5 结果",1)
    add_heading(doc,"5.1 RQ1：DRTP 改变了正式队列上的故障协同端点",2)
    add_para(doc,"正式五种子队列的固定 10M 终点评价显示，DRTP 在四个汇总端点上均高于 UTR。尤其是故障相关端点，F0、非名义平均和非名义最差回报的配对平均增益分别为 52.13、55.00 和 63.01，且五个训练种子全部呈正向。该结果表明，在统一的策略结构、训练预算和故障支持集合下，有界自适应暴露分配与更高的队列级故障任务端点相联系。名义端点的配对平均增益为 33.48，4/5 种子为正，说明这一训练干预并未以系统性降低名义任务为代价。")
    add_table(doc,["端点","UTR 均值","DRTP 均值","DRTP−UTR 配对均值","正向种子"],[
        ["名义 J","174.30","207.78","+33.48","4/5"],
        ["F0 J","144.64","196.77","+52.13","5/5"],
        ["非名义平均 J","144.70","199.70","+55.00","5/5"],
        ["非名义最差 J","124.44","187.45","+63.01","5/5"],
    ],[3.6,2.7,2.7,3.6,2.0]); add_caption(doc,"表 3 正式五种子队列的固定终点任务回报")
    doc.add_picture(str(FIG/"fig4_formal_results.png"),width=Cm(15.7)); add_caption(doc,"图 4 正式五种子队列的任务端点均值。柱上数字为 DRTP−UTR 配对平均增益。")
    add_heading(doc,"5.2 RQ2：任务收益伴随可见的可靠性权衡",2)
    add_para(doc,"非名义条件平均超时率从 UTR 的 0.8742 降至 DRTP 的 0.6938，五个种子中四个表现为下降。故障触发审计显示，两臂均实际接触故障条件，failure exposure 分别为 0.9986 和 0.9935，因此任务端点差异不能由“某一臂未被故障触发”解释。与此同时，平均碰撞率从 0.0051 变为 0.0080；两臂约束违规率均为零。因而，本文的主张是故障任务端点与平均超时的改善，而不是全面的安全性保证。")
    add_table(doc,["指标","UTR","DRTP","论文解释"],[
        ["非名义平均超时率","0.8742","0.6938","平均下降；并非逐种子单调"],
        ["非名义平均碰撞率","0.0051","0.0080","不支持全面安全改善"],
        ["约束违规率","0","0","共同为零，非 DRTP 独占优势"],
        ["故障触发暴露","0.9986","0.9935","两臂均实际经历故障"],
    ],[3.5,2.2,2.2,5.4]); add_caption(doc,"表 4 正式队列的可靠性与故障触发指标")
    add_heading(doc,"5.3 RQ3：采样遥测验证干预发生，但不替代机制识别",2)
    add_para(doc,"DRTP 训练日志记录了故障组概率、组级回报 EMA、难度代理、选择计数以及边界投影状态。与保持均匀分配的 UTR 相比，这些日志确认采样器在训练中确实产生了非均匀的故障暴露。该证据支持“DRTP 被实际执行”的解释，也与训练暴露分配是唯一主要差异的设计相一致。不过，遥测本身不能证明某个 q 变化恢复了特定信息路径、降低了某类梯度冲突，或产生了超出本实验合同的泛化机制。")
    doc.add_picture(str(FIG/"fig5_mechanism_schematic.png"),width=Cm(15.3)); add_caption(doc,"图 5 UTR 与 DRTP 的训练暴露机制示意。该图为方法结构图，不报告单一训练种子的实测最终概率。")
    add_heading(doc,"5.4 RQ4：跨队列材料限定主结论的适用范围",2)
    add_para(doc,"零训练交叉评价诊断显示，正式 cohort 在两个冻结评价带上，F0、非名义平均和非名义最差端点的配对方向均保持正向。相反，独立训练 cohort 在相同类型端点上表现为非正向方向。该结果表明：在同一正式训练 cohort 内，主要方向不依赖于单一评价带；但训练队列间的方向并不可以被假定为同质。基于这一事实，本文将正式五种子受控比较作为主要结果，而不将后续独立 cohort 合并为更大的确认性样本。PLR-style 外部比较可说明通用优先级机制在部分队列也具有竞争性，但其结构和信号与 UTR–DRTP 主要因果对照不同，不能据此替换核心比较。")
    # Discussion
    add_heading(doc,"6 讨论",1)
    add_heading(doc,"6.1 为什么训练暴露分配可能影响拓扑退化下的协同",2)
    add_para(doc,"继电故障使协作任务面对的并非统一强度的扰动，而是一组在时序、持续时间和信息支撑路径上不同的条件。均匀暴露提供了覆盖，但并未对当前训练中持续落后的条件投入额外学习机会。DRTP 将组级名义相对回报差异转化为受限的训练资源分配，从而使策略在训练期间反复接触相对困难的故障组；概率边界同时保留其他组的覆盖。正式队列中故障平均与最差端点的共同提升，与这种“覆盖加重点”的训练塑形解释一致。")
    add_heading(doc,"6.2 与通用优先级和课程组织的关系",2)
    add_para(doc,"DRTP 的价值不应被表述为通用优先级方法的普遍替代。通用优先级采样可能同样改变学习轨迹；课程学习也可通过不同任务组织改善训练稳定性。DRTP 的具体贡献在于将采样单元限制为冻结的、具有故障时序和信息路径语义的条件组，并把名义相对表现作为组级重加权信号。在本文证据范围内，这一设计提供了一个可审计且不修改执行期策略接口的拓扑退化训练方案。更强的“拓扑语义本身不可替代”主张需要 Fixed-DRTP 或 Random-DRTP 语义消融的独立结果支持。")
    add_heading(doc,"6.3 适用范围与局限",2)
    add_para(doc,"本文使用轻量三自由度仿真环境，未开展实飞验证，也未将训练支持内的故障条件写成严格 OOD。正式证据来自一个五种子受控 cohort；独立 cohort 的反向结果表明跨训练随机性的稳定性仍是需要被明确报告的边界。任务回报与超时表现改善不意味着所有安全指标同步改善，碰撞率的变化已在表 4 中单独呈现。此外，当前仓库中的早期 6-UAV 结果存在故障注入时序无效的问题，不能进入本文证据链；修复后的 6-UAV 协议、运行开销与语义消融在获得可核验结果前均不作为结论依据。")
    # Conclusion
    add_heading(doc,"7 结论",1)
    add_para(doc,"本文从通信拓扑退化改变协同信息结构这一问题出发，研究了训练阶段故障条件暴露的可控分配。在三 UAV 继电依赖协作环境中，DRTP 保持执行期信息边界、策略骨干、PPO 目标和训练预算不变，仅以有界自适应机制重分配冻结故障组的 reset 质量。正式五种子固定终点评价表明，DRTP 在 F0、非名义平均和非名义最差回报上均呈现一致的配对正向方向，并伴随较低的平均超时率。该结果支持在所评估合同内采用拓扑语义驱动的训练暴露塑形；其跨队列稳定性、安全收益、跨规模迁移和语义独立作用仍应通过后续冻结实验进一步检验。")
    # References
    add_heading(doc,"参考文献",1)
    refs=[
        "[1] 有人—无人机协同作战：概念、技术与挑战. 航空学报, 2024.（卷期页码待核验）",
        "[2] Within-visual-range air combat maneuver decision-making in obstructed environments via a curriculum self-play soft actor-critic with an attention mechanism. Defence Technology.（作者、卷期页码待核验）",
        "[3] 无人集群博弈对抗系统仿真验证及决策关键技术综述. 系统仿真学报.（卷期页码待核验）",
        "[4] 熊威. 面向有人—无人机协同打击的智能决策方法研究. 系统工程与电子技术.（卷期页码待核验）",
        "[5] A Hierarchical Deep Reinforcement Learning Framework for 6DOF Air Combat. IEEE Transactions on Systems, Man, and Cybernetics: Systems, 2023.",
        "[6] Decision-making and confrontation in close-range air combat based on reinforcement learning. Chinese Journal of Aeronautics, 2025. DOI: 10.1016/j.cja.2025.103526.",
        "[7] TODO（需补证）：DRTP 与 UTR 正式实验合同、评价带清单和结果归档的可公开引用形式。"
    ]
    for ref in refs:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.3); p.paragraph_format.first_line_indent=Cm(-.3); p.paragraph_format.line_spacing=1.2; r=p.add_run(ref); cn_font(r,"宋体",9.2)
    # footer page number
    for section in doc.sections:
        footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run("DRTP 中文投稿初稿 V2  |  "); cn_font(r,"宋体",8,color=(100,100,100));
        fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instrText=OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text='PAGE'; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); r._r.append(fldChar1); r._r.append(instrText); r._r.append(fldChar2)
    out=OUT/"DRTP_FINAL_MANUSCRIPT_SUBMISSION_READY_V2.docx"; doc.save(out); return out


if __name__ == "__main__":
    print(build_doc())
