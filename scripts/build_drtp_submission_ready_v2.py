"""Create the visual submission-ready V2 of the frozen Chinese DRTP manuscript.

All plotted numerical values originate from the aggregate values stated in the
manuscript source.  Panels without archived per-seed/q trajectories are visibly
marked as data-import slots rather than populated with synthetic observations.
"""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLUE, TEAL, ORANGE, GREY, LIGHT = "#164C78", "#087E8B", "#D97531", "#667085", "#F2F5F8"
plt.rcParams.update({"font.family": ["Arial", "Microsoft YaHei"], "font.size": 9,
                     "axes.unicode_minus": False, "svg.fonttype": "none", "pdf.fonttype": 42})


def save(fig, stem: Path):
    stem.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(stem.with_suffix(".png"), dpi=600, bbox_inches="tight", facecolor="white")
    fig.savefig(stem.with_suffix(".tiff"), dpi=600, bbox_inches="tight", facecolor="white", pil_kwargs={"compression": "tiff_lzw"})
    fig.savefig(stem.with_suffix(".svg"), bbox_inches="tight", facecolor="white")
    fig.savefig(stem.with_suffix(".pdf"), bbox_inches="tight", facecolor="white")
    plt.close(fig)


def box(ax, xy, text, color=BLUE, width=0.18, height=0.18):
    x, y = xy
    patch = FancyBboxPatch((x, y), width, height, boxstyle="round,pad=0.012,rounding_size=0.02",
                           linewidth=1.4, edgecolor=color, facecolor="white")
    ax.add_patch(patch)
    ax.text(x + width / 2, y + height / 2, text, ha="center", va="center", color=color, weight="bold", wrap=True)
    return (x + width, y + height / 2)


def arrow(ax, start, end, color=GREY):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=13, linewidth=1.25, color=color))


def fig1(path: Path):
    fig, ax = plt.subplots(figsize=(10.2, 3.0)); ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
    p1 = box(ax, (0.02, .40), "冻结拓扑\n故障空间", TEAL)
    p2 = box(ax, (0.24, .40), "名义组 N +\n故障组 F", BLUE)
    p3 = box(ax, (0.46, .40), "相对名义的\n困难度估计", ORANGE)
    p4 = box(ax, (0.68, .40), "受约束的\nq 更新", ORANGE)
    p5 = box(ax, (0.86, .40), "reset 条件\n→ MAPPO 训练", BLUE, width=.12)
    for a, b in zip((p1,p2,p3,p4), ((.24,.49),(.46,.49),(.68,.49),(.86,.49))): arrow(ax, a, b)
    ax.text(.50, .16, "保持不变：actor · critic · PPO objective · reward · observation · action interface · transition",
            ha="center", va="center", color=GREY, bbox=dict(boxstyle="round,pad=.45", fc=LIGHT, ec="#D0D5DD"))
    ax.text(.50, .88, "DRTP 仅在环境 reset 阶段分配训练暴露，不改策略学习接口", ha="center", weight="bold", color=BLUE, fontsize=11)
    save(fig, path)


def fig2(path: Path):
    fig, axes = plt.subplots(1, 3, figsize=(10.2, 3.0), sharey=True)
    data = [("UTR", "均匀条件采样", "固定支持集\n无自适应信号", BLUE),
            ("PLR-style", "通用优先分配", "priority-driven\n条件重放", ORANGE),
            ("DRTP", "拓扑语义化分配", "名义锚点 + 组级困难度\n有界 q 更新", TEAL)]
    for ax, (name, title, note, color) in zip(axes, data):
        ax.axis("off"); ax.set_xlim(0,1); ax.set_ylim(0,1)
        ax.add_patch(FancyBboxPatch((.06,.08),.88,.84, boxstyle="round,pad=.02,rounding_size=.03", fc="white", ec=color, lw=2))
        ax.text(.5,.78,name, ha="center", weight="bold", color=color, fontsize=13)
        ax.text(.5,.60,title, ha="center", color="#1D2939", weight="bold")
        ax.annotate("训练条件",(.5,.43),ha="center",bbox=dict(boxstyle="round,pad=.35",fc=LIGHT,ec="#D0D5DD"))
        ax.annotate("MAPPO",(.5,.22),ha="center",bbox=dict(boxstyle="round,pad=.35",fc="white",ec=color))
        ax.add_patch(FancyArrowPatch((.5,.38),(.5,.27),arrowstyle="-|>",mutation_scale=12,color=color))
        ax.text(.5,.97,note,ha="center",va="top",fontsize=8,color=GREY)
    fig.text(.5,.01,"机制比较图不表达性能排名；三者在匹配实验中共享环境、网络、奖励、训练预算与 PPO。",ha="center",color=GREY)
    save(fig, path)


def fig3(path: Path):
    fig, ax = plt.subplots(figsize=(8.8, 3.25)); ax.axis("off"); ax.set_xlim(0,10); ax.set_ylim(0,4)
    ax.text(5,3.55,"DRTP 采样概率 q 的可审计导入位",ha="center",weight="bold",color=BLUE,fontsize=12)
    for i, (label, color) in enumerate(zip(("F0","TE","TL","DS","DL","CP"), (BLUE,TEAL,ORANGE,"#7A5AF8","#12B76A","#EAAA08"))):
        y=2.75-i*.34; ax.plot([1,8.9],[y,y],color="#EAECF0",lw=4); ax.scatter([1],[y],s=55,color=color); ax.text(.72,y,label,ha="right",va="center",color=color,weight="bold")
    ax.axvline(1,color=GREY,ls="--",lw=1); ax.text(1,.45,"q0：均匀初始化",ha="center",color=GREY)
    ax.axvline(8.9,color=GREY,ls="--",lw=1); ax.text(8.9,.45,"终点：冻结日志导入",ha="center",color=GREY)
    ax.text(5,1.15,"该面板只接受与正式训练轨迹一一对应的 sampler CSV。\n当前论文资产未归档最终 q 日志，故不以汇总结果代替轨迹或合成曲线。",
            ha="center",va="center",bbox=dict(boxstyle="round,pad=.6",fc=LIGHT,ec="#D0D5DD"),color="#344054")
    save(fig, path)


def fig4(path: Path):
    fig, axes = plt.subplots(1,2,figsize=(8.8,3.4),sharey=True)
    vals={"A":{"UTR":(177.02,64.53,79.75),"DRTP":(216.66,23.48,191.49)},"B":{"UTR":(187.18,21.66,164.98),"DRTP":(210.34,30.54,172.03)}}
    for ax,(cohort,v) in zip(axes,vals.items()):
        names=list(v); means=[v[x][0] for x in names]; sds=[v[x][1] for x in names]; worst=[v[x][2] for x in names]
        x=np.arange(2); ax.bar(x,means,yerr=sds,capsize=5,color=[BLUE,TEAL],width=.58,alpha=.9)
        ax.scatter(x,worst,color="#101828",marker="D",s=34,zorder=4,label="最差 seed")
        ax.set_xticks(x,names); ax.set_title(f"Cohort {cohort}（n=5）",weight="bold",color=BLUE); ax.grid(axis="y",alpha=.22); ax.set_ylim(0,300)
        for xi,m in zip(x,means): ax.text(xi,m+9,f"{m:.2f}",ha="center",weight="bold",fontsize=8)
    axes[0].set_ylabel("扰动回报 J_perturbed"); axes[1].legend(loc="upper right",frameon=False,fontsize=8)
    fig.text(.5,.01,"柱：cohort 均值；误差线：样本 SD；菱形：该 cohort 的最差训练 seed。逐 seed 配对散点待冻结 per-seed CSV 导入。",ha="center",color=GREY,fontsize=8)
    fig.tight_layout(rect=(0,.07,1,1)); save(fig,path)


def fig5(path: Path):
    fig, ax = plt.subplots(figsize=(8.8,3.5))
    labels=["A\nStructural", "B\nStructural", "A\nParameter", "B\nParameter"]
    vals=[22.77,11.96,51.71,17.48]
    colors=[TEAL,TEAL,BLUE,BLUE]
    bars=ax.bar(np.arange(4),vals,color=colors,width=.62)
    ax.axhline(0,color="#344054",lw=1); ax.set_xticks(np.arange(4),labels); ax.set_ylabel("配对平均回报差（DRTP − UTR）")
    ax.set_title("冻结 OOD 条件下的 cohort-level 差值",loc="left",weight="bold",color=BLUE); ax.grid(axis="y",alpha=.22)
    for b,v in zip(bars,vals): ax.text(b.get_x()+b.get_width()/2,v+1.3,f"+{v:.2f}",ha="center",weight="bold")
    ax.text(.99,.95,"每项 n=5；结构与参数 shift 分开报告",transform=ax.transAxes,ha="right",va="top",fontsize=8,color=GREY)
    fig.tight_layout(); save(fig,path)


def fig6(path: Path):
    fig, ax=plt.subplots(figsize=(8.8,3.3)); ax.axis("off"); ax.set_xlim(0,1); ax.set_ylim(0,1)
    ax.add_patch(FancyBboxPatch((.09,.14),.82,.70,boxstyle="round,pad=.03,rounding_size=.03",fc=LIGHT,ec="#98A2B3",lw=1.4,ls="--"))
    ax.text(.5,.65,"6-UAV 跨尺度正式验证",ha="center",weight="bold",color=BLUE,fontsize=14)
    ax.text(.5,.45,"UTR / Original DRTP · 5 个 fresh training seeds · 固定 10M endpoint",ha="center",color="#344054")
    ax.text(.5,.30,"[TBD：仅在正式训练、固定终点评估与聚合完成后导入]\n将报告均值、中位数、lower tail、success、timeout、collision 与 paired delta。",ha="center",color=GREY,fontsize=9)
    save(fig,path)


FIG_CAPTIONS = {
    "图 1 DRTP 整体框架图": "fig1_framework.png",
    "图 2 UTR、PLR-style 与 DRTP 的训练暴露机制比较": "fig2_mechanism_comparison.png",
    "图 3 DRTP 训练过程中故障组采样概率": "fig3_sampling_data_slot.png",
    "图 4 A/B cohort 的 seed 级配对结果": "fig4_cohort_summary.png",
    "图 5 冻结 OOD 的配对鲁棒性结果": "fig5_ood.png",
    "图 6 六无人机跨尺度": "fig6_6uav_placeholder.png",
}

DISPLAY_CAPTIONS = {
    "图 1 DRTP 整体框架图": "图 1 DRTP 整体框架图",
    "图 2 UTR、PLR-style 与 DRTP 的训练暴露机制比较": "图 2 UTR、PLR-style 与 DRTP 的训练暴露机制比较",
    "图 3 DRTP 训练过程中故障组采样概率": "图 3 DRTP 训练过程中故障组采样概率 q 的审计导入位",
    "图 4 A/B cohort 的 seed 级配对结果": "图 4 A/B cohort 的终点汇总与 seed 级配对图位",
    "图 5 冻结 OOD 的配对鲁棒性结果": "图 5 冻结 OOD 的 cohort-level 配对鲁棒性结果",
    "图 6 六无人机跨尺度": "图 6 六无人机跨尺度验证（待正式实验完成后由最终聚合 CSV 生成）",
}


def insert_after(paragraph, text=""):
    new = paragraph._parent.add_paragraph(text)
    paragraph._p.addnext(new._p)
    return new


def style_tables(doc):
    for table in doc.tables:
        for i, cell in enumerate(table.rows[0].cells):
            for run in cell.paragraphs[0].runs:
                run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
        for row_idx, row in enumerate(table.rows[1:],1):
            if row_idx % 2 == 0:
                for cell in row.cells:
                    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"), "F5F8FA"); tcPr.append(shd)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs: r.font.size=Pt(8.1)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--input",type=Path,required=True); parser.add_argument("--output",type=Path,required=True); parser.add_argument("--figure-dir",type=Path,required=True); args=parser.parse_args()
    generators=[(fig1,"fig1_framework"),(fig2,"fig2_mechanism_comparison"),(fig3,"fig3_sampling_data_slot"),(fig4,"fig4_cohort_summary"),(fig5,"fig5_ood"),(fig6,"fig6_6uav_placeholder")]
    for fn,stem in generators: fn(args.figure_dir/stem)
    args.output.parent.mkdir(parents=True,exist_ok=True); shutil.copyfile(args.input,args.output)
    doc=Document(args.output); style_tables(doc)
    for para in list(doc.paragraphs):
        target=next((name for name in FIG_CAPTIONS if para.text.startswith(name)),None)
        if not target: continue
        para.text=DISPLAY_CAPTIONS[target]
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pic=insert_after(para); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.add_run().add_picture(str(args.figure_dir/FIG_CAPTIONS[target]),width=Cm(15.3))
    for para in list(doc.paragraphs):
        if para.text.startswith("[图 "):
            para._element.getparent().remove(para._element)
    for section in doc.sections:
        footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.text="DRTP 中文投稿主稿（V2）｜数值以冻结聚合产物与 manifest 为准"
        for run in footer.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor(102,112,133)
    doc.save(args.output)


if __name__ == "__main__": main()
