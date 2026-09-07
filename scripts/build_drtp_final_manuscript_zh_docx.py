"""Build the Chinese DRTP journal-manuscript draft as a formatted DOCX."""
from __future__ import annotations

from pathlib import Path
import re
import argparse

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs" / "drtp_final_paper_closure_20260907" / "DRTP_FINAL_MANUSCRIPT_ZH_DRAFT.md"
DEFAULT_OUTPUT = ROOT / "docs" / "drtp_final_paper_closure_20260907" / "DRTP_FINAL_MANUSCRIPT_ZH_DRAFT.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def clean_inline(value: str) -> str:
    value = value.replace("**", "").replace("`", "")
    replacements = {
        "\\(n=5\\)": "n=5",
        "\\(n\\)": "n",
        "\\(m_N=0.5\\)": "m_N=0.5",
        "\\(u_g=1/6\\)": "u_g=1/6",
        "\\(q_t\\in\\Delta^6\\)": "q_t ∈ Δ⁶",
        "\\(q_0(g)=1/6\\)": "q₀(g)=1/6",
        "\\(\\kappa=0.20\\)": "κ=0.20",
        "\\(\\beta=0.5\\)": "β=0.5",
        "\\(\\mathcal Q\\)": "Q",
        "\\(\\mathcal C\\)": "C",
        "\\(\\mathcal F\\)": "F",
        "\\(\\pi_\\theta\\)": "πθ",
        "\\(J_c(\\pi_\\theta)\\)": "Jc(πθ)",
        "\\(J_{\\mathrm{perturbed}}\\)": "Jperturbed",
        "\\(q_t\\)": "q_t",
        "\\(q_{t+1}\\)": "q_(t+1)",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = value.replace("\\mathcal{C}", "C").replace("\\mathcal{F}", "F").replace("\\mathcal Q", "Q")
    value = value.replace("\\pi_\\theta", "πθ").replace("\\bar J", "J̄").replace("\\bar d", "d̄")
    value = value.replace("\\{N\\}", "{N}").replace("\\mathcal F", "F")
    value = value.replace("\\Delta^6", "Δ⁶").replace("\\cup", "∪")
    value = value.replace("\\mathrm{", "").replace("}", "}")
    value = value.replace("\\(", "").replace("\\)", "")
    value = re.sub(r"\\text\{([^}]*)\}", r"\1", value)
    return value


def pretty_formula(value: str) -> str:
    compact = " ".join(value.split())
    if "mathcal{F}" in compact:
        return "F = {F0, TE, TL, DS, DL, CP}."
    if "frac{\\bar J_N-\\bar J_g}" in compact:
        return "d_g = min{2, max[0, (J̄_N − J̄_g) / max(|J̄_N|, 10⁻⁸)]}."
    if "tilde q_g" in compact:
        return "q̃_g ∝ q_t,g exp(d_g − d̄)."
    if "mathcal Q" in compact:
        return "Q = {q : Σ_g q_g = 1, 0.05 ≤ q_g ≤ 0.35}."
    return clean_inline(compact)


def add_text_paragraph(doc: Document, text: str, *, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.4
    run = paragraph.add_run(clean_inline(text))
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for col, text in enumerate(row):
            cell = cells[col]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if idx == 0:
                set_cell_shading(cell, "1F4E78")
            elif idx % 2 == 0:
                set_cell_shading(cell, "F3F7FA")
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or len(text) < 20 else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after = Pt(2)
            run = para.add_run(clean_inline(text))
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(8.5)
            if idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--footer", default="DRTP 中文期刊论文初稿")
    args = parser.parse_args()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 19), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    lines = args.source.read_text(encoding="utf-8").splitlines()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    math_lines: list[str] | None = None
    for raw in lines:
        line = raw.strip()
        if line == "\\[":
            flush_table()
            math_lines = []
            continue
        if math_lines is not None:
            if line == "\\]":
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(6)
                run = paragraph.add_run(pretty_formula(" ".join(math_lines)))
                run.font.name = "Cambria Math"
                run.font.size = Pt(10.5)
                math_lines = None
            else:
                math_lines.append(line)
            continue
        if not line or line.startswith("> **论文状态"):
            flush_table()
            continue
        if line.startswith("|") and line.endswith("|"):
            cells = [clean_inline(part.strip()) for part in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            table_rows.append(cells)
            continue
        flush_table()
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(16)
            paragraph.add_run(clean_inline(line[2:]))
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.add_run(clean_inline(line[3:]))
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(clean_inline(line[4:]))
        elif line.startswith("**表 ") or line.startswith("**补充表 ") or line.startswith("**算法 "):
            # Let Word flow tables naturally. Forced breaks leave large empty
            # regions when a short result paragraph precedes a table.
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(clean_inline(line))
            run.font.name = "黑体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            run.font.size = Pt(10)
        elif re.match(r"^\d+\. ", line):
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(clean_inline(re.sub(r"^\d+\. ", "", line)))
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_inline(line[2:]))
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)
        else:
            add_text_paragraph(doc, line)
    flush_table()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(args.footer)
    footer.runs[0].font.size = Pt(8)
    doc.core_properties.title = "动态鲁棒拓扑优先训练 中文期刊论文初稿"
    doc.core_properties.author = "[作者姓名]"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
